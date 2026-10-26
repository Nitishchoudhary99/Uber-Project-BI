"""Best-effort CV parsing into the profile 'fact bank'. The user always reviews and edits the result."""
import io
import re
from datetime import date

from .skills import extract_skills

SECTION_HEADS = {
    "experience": r"(work |professional )?experience|employment( history)?|work history",
    "education": r"education|academics|qualifications?",
    "skills": r"(technical |key |core )?skills|technologies|tech stack",
    "projects": r"projects?|personal projects",
    "summary": r"summary|profile|objective|about( me)?",
}
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?:\+?91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}")
MONTHS = "jan feb mar apr may jun jul aug sep oct nov dec".split()
DATE_RANGE = re.compile(
    r"((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}|\d{4})\s*(?:-|–|—|to)\s*"
    r"((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}|\d{4}|present|current|till date|now)", re.I)
YEARS_CLAIM = re.compile(r"(\d{1,2}(?:\.\d)?)\+?\s*(?:years|yrs)", re.I)
BULLET = re.compile(r"^\s*[•\-\*–▪●◦]\s*")


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


def _head(line: str):
    s = line.strip().strip(":").strip()
    if not s or len(s) > 40:
        return None
    for key, pat in SECTION_HEADS.items():
        if re.fullmatch(pat, s, re.I):
            return key
    return None


def _split_sections(lines: list[str]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {"top": []}
    cur = "top"
    for ln in lines:
        h = _head(ln)
        if h:
            cur = h
            out.setdefault(cur, [])
        else:
            out[cur].append(ln)
    return out


def _year(token: str) -> tuple[int, int]:
    m = re.search(r"(\d{4})", token)
    y = int(m.group(1)) if m else date.today().year
    mon = next((i + 1 for i, mm in enumerate(MONTHS) if token.lower().strip().startswith(mm)), 1)
    return y, mon


def _months_between(a: str, b: str) -> int:
    ya, ma = _year(a)
    if re.search(r"present|current|till|now", b, re.I):
        yb, mb = date.today().year, date.today().month
    else:
        yb, mb = _year(b)
    return max(0, (yb - ya) * 12 + (mb - ma))


def _parse_experience(lines: list[str]) -> list[dict]:
    exps, cur = [], None
    for raw in lines:
        ln = raw.strip()
        if not ln:
            continue
        m = DATE_RANGE.search(ln)
        if BULLET.match(ln):
            if cur is not None:
                cur["bullets"].append(BULLET.sub("", ln).strip())
            continue
        if m:
            head = (ln[:m.start()] + " " + ln[m.end():]).strip(" |,-–—@()")
            parts = [p.strip() for p in re.split(r"\s+(?:at|@|\||-|–|—)\s+|\s*,\s+", head) if p.strip()]
            cur = {"title": parts[0] if parts else head, "company": parts[1] if len(parts) > 1 else "",
                   "start": m.group(1), "end": m.group(2), "bullets": []}
            exps.append(cur)
        elif cur is not None and not cur["bullets"] and not cur["company"] and len(ln) < 80:
            cur["company"] = ln.strip(" |,-–—")
        elif cur is not None and cur["bullets"] and len(ln) > 25:
            cur["bullets"][-1] += " " + ln  # wrapped bullet line
        elif cur is None and len(ln) < 80:
            cur = {"title": ln, "company": "", "start": "", "end": "", "bullets": []}
            exps.append(cur)
    return [e for e in exps if e["title"]]


def _years_of_experience(exps: list[dict], text: str) -> float:
    months = sum(_months_between(e["start"], e["end"]) for e in exps if e.get("start") and e.get("end"))
    if months:
        return round(months / 12, 1)
    m = YEARS_CLAIM.search(text)
    return float(m.group(1)) if m else 0.0


def parse_resume(text: str) -> dict:
    lines = [l.rstrip() for l in text.replace("\r", "").split("\n")]
    nonblank = [l for l in lines if l.strip()]
    sec = _split_sections(nonblank)
    top = sec.get("top", [])
    name = next((l.strip() for l in top if 2 <= len(l.split()) <= 4 and not EMAIL.search(l) and not re.search(r"\d", l)), "")
    exps = _parse_experience(sec.get("experience", []))
    edu = []
    for l in sec.get("education", []):
        l = l.strip()
        if l:
            edu.append({"degree": BULLET.sub("", l), "institution": "", "year": (re.search(r"(19|20)\d{2}", l) or [""])[0] if re.search(r"(19|20)\d{2}", l) else ""})
    skill_text = "\n".join(sec.get("skills", []))
    declared = extract_skills(skill_text)
    skills = list(dict.fromkeys(declared + extract_skills(text)))
    summary = " ".join(l.strip() for l in sec.get("summary", []))[:600]
    projects = []
    for l in sec.get("projects", []):
        l = l.strip()
        if not l:
            continue
        if BULLET.match(l) and projects:
            projects[-1]["bullets"].append(BULLET.sub("", l))
        else:
            projects.append({"name": BULLET.sub("", l)[:100], "bullets": []})
    return {
        "name": name, "email": (EMAIL.search(text) or [""])[0] if EMAIL.search(text) else "",
        "phone": (PHONE.search(text) or [""])[0] if PHONE.search(text) else "",
        "headline": exps[0]["title"] if exps else "", "summary": summary,
        "years_experience": _years_of_experience(exps, text), "skills": skills,
        "experiences": exps, "education": edu, "projects": projects,
    }
